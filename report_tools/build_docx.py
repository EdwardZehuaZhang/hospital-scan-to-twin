"""Build the report and executive summary as Word documents from their Markdown sources.

python3 build_docx.py            -> report.docx, executive-summary.docx

Handles the Markdown subset the sources use: YAML front matter (title page), # and ## headings, paragraphs,
- and 1. lists, pipe tables, images with captions, **bold**, *italic* and `code`.
Format: A4, Times New Roman 11 pt, 1.5 line spacing, page numbers (NUS UROP report conventions).
"""
import os
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
FONT = "Times New Roman"


def set_font(run, size=None, bold=None, italic=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_inline(par, text, size=11, base_italic=False):
    """Add text with **bold**, *italic* and `code` spans."""
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if not tok:
            continue
        if tok.startswith("**"):
            r = par.add_run(tok[2:-2]); set_font(r, size, bold=True, italic=base_italic)
        elif tok.startswith("*"):
            r = par.add_run(tok[1:-1]); set_font(r, size, italic=not base_italic)
        elif tok.startswith("`"):
            r = par.add_run(tok[1:-1]); set_font(r, size - 1, italic=base_italic); r.font.name = "Consolas"
        else:
            r = par.add_run(tok); set_font(r, size, italic=base_italic)


def page_number_footer(section):
    p = section.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(); set_font(r, 10)
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt
        r._r.append(el)


def setup(doc):
    st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = st.paragraph_format; pf.line_spacing = 1.5; pf.space_after = Pt(6)
    for name, size in (("Heading 1", 14), ("Heading 2", 12), ("Title", 18)):
        h = doc.styles[name]; h.font.name = FONT; h.font.size = Pt(size); h.font.bold = True
        h.font.color.rgb = RGBColor(0, 0, 0)
        h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT); h.element.rPr.rFonts.set(qn("w:ascii"), FONT)
        h.element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        h.paragraph_format.space_before = Pt(12); h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.line_spacing = 1.15
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, side, Cm(2.54))
    page_number_footer(s)


def new_section(doc, landscape):
    from docx.enum.section import WD_ORIENT
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    w, h = (Cm(29.7), Cm(21.0)) if landscape else (Cm(21.0), Cm(29.7))
    sec.page_width, sec.page_height = w, h


def front_matter(md):
    m = re.match(r"---\n(.*?)\n---\n", md, re.S)
    if not m:
        return {}, md
    meta = dict(re.findall(r'^(\w+):\s*"?(.*?)"?\s*$', m.group(1), re.M))
    return meta, md[m.end():]


def title_page(doc, meta, full):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(120 if full else 0)
    r = p.add_run(meta.get("title", "")); set_font(r, 18 if full else 15, bold=True)
    if not full:
        if meta.get("author"):
            q = doc.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = q.add_run(meta["author"]); set_font(r, 11, italic=True)
        return
    for key in ("author", "course", "supervisor", "department", "term", "date"):
        if meta.get(key):
            q = doc.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if key == "author":
                q.paragraph_format.space_before = Pt(48)
            label = {"supervisor": "Supervisor: "}.get(key, "")
            r = q.add_run(label + meta[key]); set_font(r, 12, bold=(key == "author"))
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hex_); tcPr.append(sh)


def add_table(doc, lines):
    rows = [[c.strip() for c in l.strip().strip("|").split("|")] for l in lines if not re.match(r"^\|[\s\-|:]+\|$", l.strip())]
    n = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=n); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    small = 7.5 if n > 9 else (8 if n > 7 else 9.5)
    if n > 7:  # wide tables: a broader first column so labels do not wrap letter by letter
        t.autofit = False
        first = Cm(3.6); rest = Cm((15.9 - 3.6) / (n - 1))
        for row in t.rows:
            for j, cell in enumerate(row.cells):
                cell.width = first if j == 0 else rest
        for j, col in enumerate(t.columns):
            col.width = first if j == 0 else rest
    for i, r in enumerate(rows):
        for j in range(n):
            cell = t.cell(i, j); cell.text = ""
            p = cell.paragraphs[0]; p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(0)
            txt = r[j] if j < len(r) else ""
            if i == 0:
                rr = p.add_run(txt.replace("**", "")); set_font(rr, small, bold=True); shade(cell, "E8E8E8")
            else:
                add_inline(p, txt, small)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_image(doc, path, caption):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    if os.path.exists(os.path.join(HERE, path)):
        p.add_run().add_picture(os.path.join(HERE, path), width=Cm(15.9))
    else:
        r = p.add_run(f"[missing figure: {path}]"); set_font(r, 11, bold=True)
    c = doc.add_paragraph(); c.paragraph_format.line_spacing = 1.15
    m = re.match(r"(Figure \d+\.)(.*)", caption)
    if m:
        r = c.add_run(m.group(1)); set_font(r, 10, bold=True); add_inline(c, m.group(2), 10)
    else:
        add_inline(c, caption, 10)


def convert(md_name, docx_name, full):
    md = open(os.path.join(HERE, md_name), encoding="utf-8").read()
    meta, body = front_matter(md)
    doc = Document(); setup(doc); title_page(doc, meta, full)
    lines = body.split("\n"); i = 0; para = []

    def flush():
        if para:
            text = " ".join(s.strip() for s in para)
            p = doc.add_paragraph()
            if re.match(r"^\*(Table [A-Z]?\d+\..*)\*$", text):  # table caption
                p.paragraph_format.keep_with_next = True
                m = re.match(r"^\*(Table [A-Z]?\d+\.)(.*)\*$", text)
                r = p.add_run(m.group(1)); set_font(r, 10, bold=True); add_inline(p, m.group(2), 10)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                add_inline(p, text)
            para.clear()

    in_refs = False
    while i < len(lines):
        l = lines[i]
        if l.startswith("# "):
            flush(); title = l[2:].strip()
            if full and title.startswith("Appendix A"):
                new_section(doc, landscape=True)
            elif full and title.startswith("Appendix B"):
                new_section(doc, landscape=False)
            elif full and (title.startswith("1. ") or title == "References"):
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            doc.add_heading(title, level=1); in_refs = title == "References"
        elif l.startswith("## "):
            flush(); doc.add_heading(l[3:].strip(), level=2)
        elif l.startswith("!["):
            flush(); m = re.match(r"!\[(.*)\]\((.*)\)", l); add_image(doc, m.group(2), m.group(1))
        elif l.startswith("|"):
            flush(); tl = []
            while i < len(lines) and lines[i].startswith("|"):
                tl.append(lines[i]); i += 1
            add_table(doc, tl); continue
        elif re.match(r"^(\s*)(- |\d+\. )", l):
            flush(); m = re.match(r"^(\s*)(- |\d+\. )(.*)", l)
            item = [m.group(3)]; i += 1
            while i < len(lines) and lines[i].startswith("  ") and not re.match(r"^\s*(- |\d+\. )", lines[i]):
                item.append(lines[i].strip()); i += 1
            if m.group(2) == "- ":
                p = doc.add_paragraph(style="List Bullet")
            else:  # numbered items as plain paragraphs, so numbering restarts in every list
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.9); p.paragraph_format.first_line_indent = Cm(-0.6)
                r = p.add_run(m.group(2)); set_font(r, 11)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, " ".join(item)); continue
        elif not l.strip():
            flush()
            if in_refs:
                pass
        else:
            para.append(l)
            if in_refs:  # references: one paragraph each, hanging indent
                flush(); p = doc.paragraphs[-1]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.left_indent = Cm(1.27); p.paragraph_format.first_line_indent = Cm(-1.27)
                p.paragraph_format.line_spacing = 1.15
        i += 1
    flush()
    doc.save(os.path.join(HERE, docx_name)); print("wrote", docx_name)


if __name__ == "__main__":
    convert("report.md", "report.docx", full=True)
    convert("executive-summary.md", "executive-summary.docx", full=False)
